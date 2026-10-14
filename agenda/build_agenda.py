"""Parse World Aviation Festival agenda text (copied from hub.terrapinn.com) into Excel.

Usage: python build_agenda.py   -> reads raw/day*.txt, writes WAF_Agenda.xlsx
"""
import re
from pathlib import Path
from datetime import time

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

HERE = Path(__file__).parent
SOURCE_URL = "https://hub.terrapinn.com/a0AN200000t9wgDMAQPA06066985/agenda"

CATEGORIES = {
    "Networking", "Airline Keynote", "Airport Keynote", "Loyalty", "Sustainability",
    "AI / ML", "Customer Experience", "IROPS", "Digital & IT", "Airport - Operations",
    "Airport & Airline Collaboration", "Airport - Traveller", "Retail", "Payments",
    "Startup Theatre", "Ancillaries",
}
TIME_RE = re.compile(r"^(\d{1,2}):(\d{2})$")
RANGE_RE = re.compile(r"^(\d{1,2}:\d{2})\s*[–-]\s*(\d{1,2}:\d{2})$")
LOC_RE = re.compile(r"^(Track\s?\d|Keynote (Stage|Theatre)|Startup Theatre$|Aviation Connect Zone|Roundtable area)")
MORE_RE = re.compile(r"^\+(\d+) more speakers?$")
REG_RE = re.compile(r"^(\d+)/(\d+) registered$")
STATUS_LINES = {"Register", "Session Full", "Registered", "Leave table"}
SKIP_LINES = {"Recommended", "Show more", "Show less"}
# Company names that themselves contain a comma
COMMA_COMPANIES = ["Edgar, Dunn & Company"]
CORP_SUFFIX_RE = re.compile(r"^(Inc\.?|LLC|Ltd\.?|Limited|S\.A\.S)(\s.*)?$")
FORMAT_KEYWORDS = [
    ("interview", "CEO Interview"), ("fireside", "Fireside Chat"), ("panel", "Panel"),
    ("presentation", "Presentation"), ("welcome", "Welcome"), ("remarks", "Chair Remarks"),
    ("aviation connect", "Networking"),
]


def to_time(s):
    h, m = s.split(":")
    return time(int(h), int(m))


def split_speaker(line):
    line = line.strip().rstrip("`").strip()
    protected = line
    for i, c in enumerate(COMMA_COMPANIES):
        protected = protected.replace(c, f"@@{i}@@")
    parts = [p.strip() for p in protected.split(",")]
    merged = []
    for p in parts:
        if not p:
            continue
        if merged and CORP_SUFFIX_RE.match(p):
            merged[-1] = f"{merged[-1]}, {p}"
        else:
            merged.append(p)
    restore = lambda s: re.sub(r"@@(\d+)@@", lambda m: COMMA_COMPANIES[int(m.group(1))], s)
    merged = [restore(p) for p in merged]
    name = merged[0]
    if len(merged) == 1:
        return name, "", ""
    if len(merged) == 2:
        return name, "", merged[1]
    return name, ", ".join(merged[1:-1]), merged[-1]


def normalise_room(raw):
    if not raw:
        return "", ""
    r = raw.replace("Pavillion", "Pavilion").replace("Pav 1", "Pavilion 1").replace("Pav 2", "Pavilion 2")
    r = re.sub(r"Track\s?(\d+)", r"Track \1", r)
    r = re.sub(r"^(Track \d+)\s*,\s*(Pavilion \d)$", r"\1 (\2)", r)
    r = re.sub(r"^(Track \d+)\s*\((Pavilion \d)\)$", r"\1 (\2)", r)
    r = re.sub(r"^Aviation Connect Zone,?\s*(Pavilion \d)$", r"Aviation Connect Zone (\1)", r)
    if r in ("Keynote Stage, Pavilion 1", "Keynote Theatre (Pavilion 1)"):
        r = "Keynote Stage 1 (Pavilion 1)"
    r = re.sub(r"^Roundtable area,\s*(Pavilion \d)$", r"Roundtable Area (\1)", r)
    pav = re.search(r"Pavilion \d", r)
    return r, pav.group(0) if pav else ""


def session_format(category, title):
    head = title.split(":")[0].lower() if ":" in title else title.lower()
    for kw, label in FORMAT_KEYWORDS:
        if kw in head:
            return label
    if category == "Startup Theatre":
        return "Startup Pitch"
    return "Talk"


def parse_day(text, day_label):
    lines = [l.strip() for l in text.splitlines()]
    lines = [l for l in lines if l and l not in SKIP_LINES]
    sessions, cur_time, cur, mode = [], None, None, "expect_cat"
    i = 0

    def new_session(cat):
        s = dict(day=day_label, start=cur_time, end=None, category=cat, title="", speakers=[],
                 hidden=0, room_raw="", description=[], registered=None, capacity=None,
                 status="", kind="Session")
        sessions.append(s)
        return s

    while i < len(lines):
        line = lines[i]
        m_range = RANGE_RE.match(line)
        if m_range and i + 1 < len(lines) and lines[i + 1] in ("Roundtables", "Workshops"):
            kind = "Roundtable" if lines[i + 1] == "Roundtables" else "Workshop"
            start, end = m_range.groups()
            i += 2
            if i < len(lines) and RANGE_RE.match(lines[i]):
                i += 1
            room = lines[i]
            i += 1
            buf = []
            while i < len(lines) and not TIME_RE.match(lines[i]) and not RANGE_RE.match(lines[i]):
                l = lines[i]
                m_reg = REG_RE.match(l)
                if m_reg:
                    title, hosts = buf[-2], buf[-1]
                    s = new_session(kind)
                    s.update(start=to_time(start), end=to_time(end), title=title, room_raw=room, kind=kind,
                             registered=int(m_reg.group(1)), capacity=int(m_reg.group(2)))
                    seen = []
                    for h in hosts.split(","):
                        h = h.strip()
                        if h and h not in seen:
                            seen.append(h)
                    s["speakers"] = [(h, "", "") for h in seen]
                    statuses = []
                    while i + 1 < len(lines) and lines[i + 1] in STATUS_LINES:
                        i += 1
                        statuses.append(lines[i])
                    s["status"] = "Full" if "Session Full" in statuses else "Open"
                    s["you_registered"] = "Registered" in statuses
                    m_tab = re.match(r"^Table\s?0*(\d+):\s*", title)
                    s["table_no"] = int(m_tab.group(1)) if m_tab else None
                    buf = []
                else:
                    buf.append(l)
                i += 1
            cur, mode = None, "expect_cat"
            continue

        m_time = TIME_RE.match(line)
        if m_time:
            cur_time, cur, mode = to_time(line), None, "expect_cat"
        elif mode == "expect_title":
            cur["title"], mode = line, "speakers"
        elif mode == "speakers" and LOC_RE.match(line):
            cur["room_raw"], mode = line, "after_location"
        elif mode == "speakers" and MORE_RE.match(line):
            cur["hidden"] = int(MORE_RE.match(line).group(1))
        elif line in CATEGORIES and mode in ("expect_cat", "speakers", "after_location"):
            cur, mode = new_session(line), "expect_title"
        elif mode == "speakers":
            cur["speakers"].append(split_speaker(line))
        elif mode == "after_location":
            cur["description"].append(line)
        else:
            raise ValueError(f"{day_label}: unexpected line {i}: {line!r} (mode={mode})")
        i += 1

    for s in sessions:
        s["room"], s["pavilion"] = normalise_room(s["room_raw"])
        s["format"] = s["kind"] if s["kind"] != "Session" else session_format(s["category"], s["title"])
        desc = " ".join(s["description"])
        s["description"] = "\n".join("• " + p.strip() if n else p.strip()
                                     for n, p in enumerate(desc.split(" • "))) if desc else ""
    return sessions


# ---------------- workbook ----------------
FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
BODY_FONT = Font(name=FONT, size=10)
TITLE_FONT = Font(name=FONT, bold=True, size=14, color="1F3864")
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Side(style="thin", color="D9D9D9")


def write_table(ws, headers, rows, widths, name, time_cols=(), start_row=1):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=c, value=h)
        cell.font, cell.fill = HEAD_FONT, HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for r, row in enumerate(rows, start_row + 1):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font, cell.alignment = BODY_FONT, WRAP
            cell.border = Border(bottom=THIN)
            if c in time_cols:
                cell.number_format = "hh:mm"
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    last = f"{get_column_letter(len(headers))}{start_row + max(len(rows), 1)}"
    tab = Table(displayName=name, ref=f"A{start_row}:{last}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(tab)
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    ws.row_dimensions[start_row].height = 30


SESSION_HEADERS = ["#", "Day", "Start", "End", "Stream / Category", "Format", "Session Title",
                   "Speakers (Name – Job Title, Company)", "Speaker Names", "Organisations",
                   "Speakers Listed", "Additional Speakers (not shown on site)",
                   "Room / Stage", "Pavilion", "Room (as listed on site)", "Description / Discussion Points",
                   "Registered", "Capacity", "Places Left", "Registration Status"]
SESSION_WIDTHS = [5, 7, 7, 7, 20, 14, 50, 55, 35, 35, 9, 12, 26, 11, 26, 80, 10, 9, 9, 13]


def session_row(n, s, excel_row):
    spk = s["speakers"]
    full = "\n".join(f"{nm} – {t}, {co}" if t else (f"{nm}, {co}" if co else nm) for nm, t, co in spk)
    names = ", ".join(nm for nm, _, _ in spk)
    orgs = ", ".join(dict.fromkeys(co for _, _, co in spk if co))
    places = f"=R{excel_row}-Q{excel_row}" if s["capacity"] is not None else None
    status = s["status"]
    if s.get("you_registered"):
        status = "Open (you are registered)"
    return [n, s["day"], s["start"], s["end"], s["category"], s["format"], s["title"], full, names, orgs,
            len(spk), s["hidden"] or None, s["room"], s["pavilion"], s["room_raw"], s["description"],
            s["registered"], s["capacity"], places, status or None]


def build():
    days = {}
    for f in sorted((HERE / "raw").glob("day*.txt")):
        label = "Day " + re.search(r"\d+", f.stem).group(0)
        days[label] = parse_day(f.read_text(encoding="utf-8"), label)

    wb = Workbook()
    ov = wb.active
    ov.title = "Overview"

    # Per-day sheets + combined
    all_rows = []
    for label, sessions in days.items():
        ws = wb.create_sheet(label)
        rows = [session_row(n, s, n + 1) for n, s in enumerate(sessions, 1)]
        write_table(ws, SESSION_HEADERS, rows, SESSION_WIDTHS, label.replace(" ", ""), time_cols=(3, 4))
        all_rows.extend(sessions)
    ws = wb.create_sheet("All Sessions")
    rows = [session_row(n, s, n + 1) for n, s in enumerate(all_rows, 1)]
    write_table(ws, SESSION_HEADERS, rows, SESSION_WIDTHS, "AllSessions", time_cols=(3, 4))

    # Speaker appearances
    ws = wb.create_sheet("Speaker Appearances")
    sp_rows = []
    for s in all_rows:
        role = "Host" if s["kind"] in ("Roundtable", "Workshop") else "Speaker"
        for nm, t, co in s["speakers"]:
            sp_rows.append([nm, t, co, role, s["day"], s["start"], s["category"], s["format"], s["title"], s["room"]])
    write_table(ws, ["Name", "Job Title", "Organisation", "Role", "Day", "Start", "Stream / Category",
                     "Format", "Session Title", "Room / Stage"], sp_rows,
                [28, 45, 35, 9, 7, 7, 22, 14, 60, 26], "SpeakerAppearances", time_cols=(6,))

    # Unique speaker directory with formula-driven session counts
    ws = wb.create_sheet("Speaker Directory")
    directory = {}
    for nm, t, co, *_ in sp_rows:
        d = directory.setdefault(nm, {"titles": [], "orgs": []})
        if t and t not in d["titles"]:
            d["titles"].append(t)
        if co and co not in d["orgs"]:
            d["orgs"].append(co)
    n_app = len(sp_rows) + 1
    dir_rows = []
    for r, (nm, d) in enumerate(sorted(directory.items(), key=lambda kv: kv[0].lower()), 2):
        dir_rows.append([nm, " / ".join(d["titles"]), " / ".join(d["orgs"]),
                         f"=COUNTIF('Speaker Appearances'!$A$2:$A${n_app},A{r})"])
    write_table(ws, ["Name", "Job Title", "Organisation", "No. of Sessions"], dir_rows,
                [30, 55, 40, 12], "SpeakerDirectory")

    # Overview sheet
    ov["A1"] = "World Aviation Festival – Agenda"
    ov["A1"].font = TITLE_FONT
    info = [
        ("Source", SOURCE_URL),
        ("Days included", ", ".join(days) or "—"),
        ("Note", "Built from agenda text copied from the event hub. The site could not be fetched directly from this "
                 "environment. Speaker job titles/organisations are split from 'Name, Title, Company' on the site; "
                 "'Additional speakers' are names hidden behind '+N more' on the site and are not listed here. "
                 "Registration numbers are a snapshot at the time of copying."),
    ]
    for r, (k, v) in enumerate(info, 3):
        ov.cell(row=r, column=1, value=k).font = Font(name=FONT, bold=True, size=10)
        c = ov.cell(row=r, column=2, value=v)
        c.font, c.alignment = BODY_FONT, WRAP
    ov.row_dimensions[5].height = 110

    r0 = 8
    ov.cell(row=r0, column=1, value="Sessions by stream").font = Font(name=FONT, bold=True, size=12, color="1F3864")
    cats = sorted({s["category"] for s in all_rows})
    day_labels = list(days)
    hdr = ["Stream / Category"] + day_labels + ["Total"]
    stats = []
    n_all = len(all_rows) + 1
    for k, cat in enumerate(cats):
        er = r0 + 2 + k
        row = [cat]
        for j, dl in enumerate(day_labels):
            row.append(f"=COUNTIFS('All Sessions'!$E$2:$E${n_all},$A{er},'All Sessions'!$B$2:$B${n_all},\"{dl}\")")
        row.append(f"=SUM(B{er}:{get_column_letter(1 + len(day_labels))}{er})")
        stats.append(row)
    tot_r = r0 + 2 + len(cats)
    write_table(ov, hdr, stats, [34] + [10] * len(day_labels) + [10], "StreamSummary", start_row=r0 + 1)
    ov.cell(row=tot_r, column=1, value="Total").font = Font(name=FONT, bold=True, size=10)
    for j in range(2, len(hdr) + 1):
        col = get_column_letter(j)
        c = ov.cell(row=tot_r, column=j, value=f"=SUM({col}{r0 + 2}:{col}{tot_r - 1})")
        c.font = Font(name=FONT, bold=True, size=10)
    ov.column_dimensions["B"].width = max(ov.column_dimensions["B"].width or 10, 12)
    ov.freeze_panes = None
    # widen B for the info block without breaking the summary table look
    ov.column_dimensions["B"].width = 12
    for r in range(3, 6):
        ov.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)

    out = HERE / "WAF_Agenda.xlsx"
    wb.save(out)
    return out, days


if __name__ == "__main__":
    out, days = build()
    for d, ss in days.items():
        print(d, len(ss), "rows")
    print("wrote", out)
